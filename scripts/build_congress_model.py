#!/usr/bin/env python3
"""Build the economic model of a Latin congress as an Excel file with live formulas.

Run:  python scripts/build_congress_model.py
Writes analysis/congress_model.xlsx and prints the results that the same model gives in Python
(used to check the spreadsheet after recalculation). Inputs come from:
- data/economics/cost_benchmarks.csv (public cost benchmarks, round AA3)
- data/prices/latin_events_detail.csv (pass prices of 149 Latin events, rounds AA1 and AB)
- assumptions written below (marked as such in the file) and values to fill from organiser talks.
Needs openpyxl.
"""

import csv
import os
import statistics

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "analysis", "congress_model.xlsx")
BENCH = os.path.join(ROOT, "data", "economics", "cost_benchmarks.csv")
PRICES = os.path.join(ROOT, "data", "prices", "latin_events_detail.csv")

FILL = {
    "source": PatternFill("solid", fgColor="DDEBF7"),
    "market data": PatternFill("solid", fgColor="E2F0D9"),
    "assumption": PatternFill("solid", fgColor="FFF2CC"),
    "to fill": PatternFill("solid", fgColor="FCE4D6"),
    "formula": PatternFill("solid", fgColor="F2F2F2"),
    "head": PatternFill("solid", fgColor="0F1A3D"),
    "band": PatternFill("solid", fgColor="D6123F"),
}
BOLD = Font(bold=True)
WHITE = Font(bold=True, color="FFFFFF")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))
WRAP = Alignment(wrap_text=True, vertical="top")
EUR = '#,##0" €"'
EUR2 = '#,##0.00" €"'
PCT = "0.0%"

# ----------------------------------------------------------------------------- data on file

def market_prices():
    rows = [r for r in csv.DictReader(open(PRICES, encoding="utf-8")) if "out of scope" not in r["notes"].lower()]

    def num(v):
        try:
            return float(v)
        except ValueError:
            return None

    full = [num(r["full_pass_eur"]) for r in rows if r["sample"] == "random" and num(r["full_pass_eur"])]
    party = [num(r["party_pass_eur"]) for r in rows if num(r["party_pass_eur"])]
    by_country = {}
    for r in rows:
        v = num(r["full_pass_eur"])
        if v:
            by_country.setdefault(r["country_iso2"], []).append(v)
    by_days = {}
    for r in rows:
        v, d = num(r["full_pass_eur"]), num(r["days"])
        if v and d:
            by_days.setdefault("1-2 days" if d <= 2 else "3 days" if d == 3 else "4+ days", []).append(v)
    return full, party, by_country, by_days


FULL, PARTY, BY_COUNTRY, BY_DAYS = market_prices()
PASS_MEDIAN = statistics.median(FULL)
PARTY_MEDIAN = statistics.median(PARTY)

# Countries: VAT standard (PwC, AA3), a reduced rate that MAY apply to cultural events (to verify),
# Eventbrite fee (AA3), music rights per night (AA3: GEMA DE, AKM AT; others not found),
# price level index 2024 (Eurostat tec00120, data/cities/city_indicators.csv).
COUNTRIES = [
    # code, name, vat_std, vat_reduced, platform_pct, platform_fixed, music_night, pli
    ("AT", "Austria", 0.20, 0.13, 0.055, 0.99, 54.33, 111.7),
    ("DE", "Germany", 0.19, 0.07, 0.055, 0.99, 152.0, 108.8),
    ("ES", "Spain", 0.21, 0.10, 0.09, 0.19, None, 91.3),
    ("IT", "Italy", 0.22, 0.10, 0.055, 0.99, None, 97.0),
    ("FR", "France", 0.20, 0.055, 0.055, 0.99, None, 111.2),
    ("NL", "Netherlands", 0.21, 0.09, 0.04, 0.99, None, 115.1),
    ("PL", "Poland", 0.23, 0.08, 0.025, 1.37, None, 72.3),
    ("CZ", "Czechia", 0.21, 0.12, 0.025, 0.78, None, 88.5),
    ("GB", "United Kingdom", 0.20, 0.20, 0.0695, 0.70, None, 125.1),
]
BASE_PLI = 108.8  # most cost benchmarks are German prices

# Global inputs: key, label, unit, value, basis, note
GLOBAL = [
    ("country", "Country of the event", "code", "AT", "assumption", "Choose from the list; drives VAT, ticket fees, music rights and the local price level"),
    ("reduced_vat", "Use the reduced VAT rate?", "yes/no", "no", "assumption", "Dance parties are often taxed at the standard rate; workshops and cultural shows may get a reduced rate. Check with a tax adviser"),
    ("scale_prices", "Scale local costs to the country's price level?", "yes/no", "yes", "assumption", "Most cost benchmarks are German prices; 'yes' multiplies local costs by price level of the country / Germany"),
    ("days", "Event days", "days", 3, "market data", "Median length of 149 Latin events (rounds AA1, AB)"),
    ("nights", "Party nights", "nights", 3, "market data", "Typical: Friday to Sunday"),
    ("pass_price", "Full pass, listed price", "EUR incl. VAT", round(PASS_MEDIAN), "market data", f"Median of {len(FULL)} randomly drawn Latin events (AA1, AB). Country medians on the sheet Market"),
    ("discount", "Average discount vs listed price", "%", 0.10, "assumption", "Early birds, group and promo codes"),
    ("party_price", "Party-only pass", "EUR incl. VAT", round(PARTY_MEDIAN), "market data", f"Median of {len(PARTY)} events that publish a party pass"),
    ("hotel_share", "Dancers booking the organiser's hotel package", "%", 0.30, "assumption", "To check with organisers"),
    ("hotel_pppn", "Hotel package price per person per night", "EUR", 60, "to fill", "Room price agreed with the hotel"),
    ("hotel_comm", "Hotel commission paid to the organiser", "%", 0.12, "source", "Kempinski group terms: 12% on net room revenue (AA3, K14)"),
    ("artist_fee", "Fee per teaching artist (couple or solo), weekend", "EUR", 800, "to fill", "Not public. Placeholder only: ask organisers and artists"),
    ("people_per_unit", "People per teaching artist (couples count 2)", "people", 1.8, "assumption", "Most line-ups are couples"),
    ("fly_share", "Share of artists who fly in", "%", 0.80, "assumption", "27 of the 51 most booked artists with a stated base live in Spain (AB)"),
    ("flight", "Average one-way fare", "EUR", 50.66, "source", "Ryanair average fare FY26, without extras (AA3, K10)"),
    ("flight_extra", "Extras on flights (bags, dates)", "%", 0.30, "assumption", ""),
    ("artist_room", "Hotel room per artist per night", "EUR", 90, "assumption", "One room per couple or solo"),
    ("per_diem", "Meals per artist person per day", "EUR", 30, "assumption", ""),
    ("dj_fee", "DJ fee per night", "EUR", 250, "to fill", "Not public"),
    ("dj_travel", "Travel per DJ", "EUR", 150, "assumption", ""),
    ("room_day", "Workshop room, rent per day", "EUR", 300, "assumption", "Public venue prices found only for small halls (Graz EUR 400-937 per day, about 240 people)"),
    ("floor_needed", "Rent a dance floor for the party hall?", "yes/no", "yes", "assumption", "'no' if the venue has a suitable floor"),
    ("floor_m2", "Dance floor, price per m2 for the event", "EUR", 15, "source", "Parquet floor 36 m2 for EUR 540 per day (AA3, K6); assumed charged once for a weekend hire"),
    ("sound_room", "Sound, workshop room, per day", "EUR", 420, "source", "Package up to 100 people (AA3, K7)"),
    ("security_rate", "Security staff, per hour", "EUR", 46.84, "source", "German security collective agreement (AA3, K9)"),
    ("hours_night", "Hours of security and first aid per night", "hours", 7, "assumption", ""),
    ("aid_rate", "First aid staff, per hour", "EUR", 40, "assumption", "No public tariff found"),
    ("photo_day", "Photographer or videographer, per day", "EUR", 900, "source", "Full day in ES and DE (AA3, K13)"),
    ("wristband", "Wristband or printed pass, per dancer", "EUR", 0.10, "source", "Printed vinyl wristbands, EUR 100 per 1,000 (AA3, K12)"),
    ("meals_vol", "Meals per volunteer per day", "EUR", 20, "assumption", "Volunteers get a free pass and meals"),
    ("pay_pct", "Card payment fee", "%", 0.015, "source", "Stripe, EEA cards (AA3, K3)"),
    ("pay_fixed", "Card payment fee, fixed per ticket", "EUR", 0.25, "source", "Stripe, EEA cards (AA3, K3)"),
    ("marketing_pct", "Marketing (ads, design, promotion)", "% of revenue", 0.05, "assumption", ""),
    ("contingency", "Contingency on production costs", "%", 0.10, "assumption", ""),
]

# Scenario inputs: key, label, unit, (small, medium, large), basis, note
SCEN = [
    ("dancers", "Paying dancers", "dancers", (300, 800, 2000), "assumption", "Three sizes; the 10 largest European congresses average about 2,400 (round Z)"),
    ("full_share", "Share of dancers on a full pass", "%", (0.75, 0.75, 0.75), "assumption", "The rest buy party-only passes"),
    ("units", "Teaching artists (couples or solo)", "artists", (6, 12, 25), "market data", "Line-ups have a median of 9 teaching artists (61 events, AA1, AB)"),
    ("djs", "DJs", "DJs", (2, 4, 6), "assumption", ""),
    ("rooms", "Workshop rooms", "rooms", (2, 3, 5), "assumption", ""),
    ("hall_day", "Party hall, rent per day", "EUR", (1000, 2500, 6000), "to fill", "Large halls rarely publish prices; congress hotels often waive rent against the room block"),
    ("floor_area", "Dance floor to rent", "m2", (100, 200, 400), "assumption", "About 0.2-0.35 m2 per dancer: not everyone dances at once"),
    ("sound_main", "Sound and light, party hall(s), per day", "EUR", (545, 985, 1970), "source", "Packages up to 400 and up to 1,000 people; large = two halls (AA3, K7)"),
    ("music_factor", "Music rights size factor", "x", (1, 1.5, 3), "assumption", "GEMA example is for about 500 m2 and 500-1,500 people"),
    ("guards", "Security staff per night", "people", (2, 3, 6), "assumption", ""),
    ("aid_staff", "First aid staff per night", "people", (1, 2, 3), "assumption", ""),
    ("volunteers", "Volunteers", "people", (8, 16, 35), "assumption", ""),
    ("photographers", "Photographers and videographers", "people", (1, 2, 2), "assumption", ""),
    ("insurance", "Event liability insurance", "EUR per event", (200, 400, 994), "source", "UK table: about EUR 82-994 by attendance (AA3, K8)"),
    ("deco", "Decoration and stage", "EUR", (300, 800, 2000), "assumption", ""),
    ("admin", "Website, software, admin", "EUR", (500, 1000, 2000), "assumption", ""),
    ("legal", "Accounting and legal", "EUR", (500, 800, 1500), "assumption", ""),
    ("sponsor", "Sponsors and stands (income)", "EUR", (0, 2000, 8000), "to fill", ""),
    ("bar", "Share of bar sales (income)", "EUR", (0, 0, 0), "to fill", "Depends on the venue deal"),
]

COLS = ["C", "D", "E"]

# ----------------------------------------------------------------------------- python model (check)

def py_model():
    g = {k: v for k, _, _, v, _, _ in GLOBAL}
    c = {row[0]: row for row in COUNTRIES}[g["country"]]
    vat = c[3] if g["reduced_vat"] == "yes" else c[2]
    music = c[6] if c[6] is not None else 152.0
    f = c[7] / BASE_PLI if g["scale_prices"] == "yes" else 1.0
    out = []
    for i in range(3):
        s = {k: v[i] for k, _, _, v, _, _ in SCEN}
        d = s["dancers"]
        gross = d * (s["full_share"] * g["pass_price"] + (1 - s["full_share"]) * g["party_price"]) * (1 - g["discount"])
        net = gross / (1 + vat)
        hotel = d * g["hotel_share"] * g["hotel_pppn"] * g["nights"] * g["hotel_comm"]
        rev = net + hotel + s["sponsor"] + s["bar"]
        lines = {
            "Artist fees": s["units"] * g["artist_fee"],
            "Artist flights": s["units"] * g["people_per_unit"] * g["fly_share"] * g["flight"] * 2 * (1 + g["flight_extra"]),
            "Artist hotel": s["units"] * g["artist_room"] * g["days"] * f,
            "Artist meals": s["units"] * g["people_per_unit"] * g["per_diem"] * g["days"] * f,
            "DJs": s["djs"] * (g["dj_fee"] * g["nights"] + g["dj_travel"]),
            "Party hall": s["hall_day"] * g["days"] * f,
            "Workshop rooms": s["rooms"] * g["room_day"] * g["days"] * f,
            "Dance floor": (s["floor_area"] * g["floor_m2"] * f) if g["floor_needed"] == "yes" else 0,
            "Sound and light": (s["sound_main"] + s["rooms"] * g["sound_room"]) * g["days"] * f,
            "Music rights": music * s["music_factor"] * g["nights"],
            "Decoration and stage": s["deco"] * f,
            "Security": s["guards"] * g["hours_night"] * g["security_rate"] * g["nights"] * f,
            "First aid": s["aid_staff"] * g["hours_night"] * g["aid_rate"] * g["nights"] * f,
            "Photo and video": s["photographers"] * g["photo_day"] * g["days"] * f,
            "Volunteers' meals": s["volunteers"] * g["meals_vol"] * g["days"] * f,
            "Wristbands and passes": d * g["wristband"],
            "Insurance": s["insurance"],
            "Website, software, admin": s["admin"],
            "Accounting and legal": s["legal"],
        }
        prod = sum(lines.values())
        lines["Contingency"] = prod * g["contingency"]
        lines["Ticket platform fees"] = gross * c[4] + d * c[5]
        lines["Card payment fees"] = gross * g["pay_pct"] + d * g["pay_fixed"]
        lines["Marketing"] = rev * g["marketing_pct"]
        cost = sum(lines.values())
        out.append((rev, cost, rev - cost, lines))
    return out

# ----------------------------------------------------------------------------- workbook

def style_header(ws, row, values, widths=None):
    for j, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=j, value=v)
        cell.font = WHITE
        cell.fill = FILL["head"]
        cell.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + j)].width = w


def build():
    wb = Workbook()

    # README -------------------------------------------------------------
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 110
    text = [
        ("Latin congress model", True),
        ("What it does: estimates revenue, costs, result and break-even of a salsa and bachata congress in three sizes (300, 800 and 2,000 paying dancers). It is a planning tool, not a forecast.", False),
        ("", False),
        ("How to use it", True),
        ("1. Go to Inputs. Change only coloured cells. Choose the country in the first row: VAT, ticket fees, music rights and the local price level follow.", False),
        ("2. Orange cells are values nobody publishes (artist fees, hall rent, hotel price, sponsors): replace the placeholders with what organisers, venues and artists tell you.", False),
        ("3. Read the results on Results; Model shows every line; Sensitivity shows which lines move the result most.", False),
        ("", False),
        ("Colours", True),
        ("Blue = public source (price list, tariff, fee page; see Sources).", False),
        ("Green = market data from our sample of Latin events (see Market).", False),
        ("Yellow = assumption, reasonable but not measured.", False),
        ("Orange = to fill: no public value; the number shown is a placeholder.", False),
        ("Grey = formula, do not type here.", False),
        ("", False),
        ("Sheets", True),
        ("Inputs: all inputs, with unit, basis and note. Country: country parameters. Model: revenue and cost lines per scenario. Results: summary and break-even. Sensitivity: effect of a 20% change of each line (medium scenario). Market: pass prices of 149 Latin events. Sources: public cost benchmarks with links.", False),
        ("", False),
        ("Definitions", True),
        ("Prices paid by dancers include VAT; revenue in the model is net of VAT. Break-even dancers = dancers needed to cover the costs of the chosen set-up (artists, rooms, staff stay as in that scenario). Break-even pass price = full pass price that gives a result of zero with the scenario's dancers.", False),
        ("", False),
        ("Limits", True),
        ("Artist fees, hall rent and hotel prices are not public: until they are filled, the result is only as good as the placeholders. VAT: standard rates by default; reduced cultural rates (column on Country) must be checked per case. Music rights are known only for Germany and Austria; other countries use the German value. The model has no Red Bull or other partner module.", False),
        ("", False),
        ("Built by scripts/build_congress_model.py from data/economics/cost_benchmarks.csv and data/prices/latin_events_detail.csv (Dance Scene Atlas project, October 2026).", False),
    ]
    for i, (t, b) in enumerate(text, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.alignment = WRAP
        if b:
            c.font = Font(bold=True, size=14 if i == 1 else 12)
    for key, row in (("source", 10), ("market data", 11), ("assumption", 12), ("to fill", 13), ("formula", 14)):
        ws.cell(row=row, column=1).fill = FILL[key]

    # Country ------------------------------------------------------------
    wc = wb.create_sheet("Country")
    style_header(wc, 1, ["Code", "Country", "VAT standard", "VAT reduced (may apply, verify)", "Ticket platform fee %", "Ticket platform fee, fixed EUR", "Music rights per night EUR", "Price level (EU = 100)", "Market median full pass EUR", "Events with a price"],
                 [8, 16, 12, 16, 14, 14, 16, 14, 16, 12])
    for i, (code, name, vs, vr, pp, pf, mu, pli) in enumerate(COUNTRIES, 2):
        prices = BY_COUNTRY.get(code, [])
        vals = [code, name, vs, vr, pp, pf, mu if mu is not None else "", pli,
                round(statistics.median(prices), 1) if len(prices) >= 5 else "", len(prices)]
        for j, v in enumerate(vals, 1):
            cell = wc.cell(row=i, column=j, value=v)
            cell.border = THIN
            if j in (3, 4, 5):
                cell.number_format = PCT
            if j in (3, 5, 6, 7, 8):
                cell.fill = FILL["source"]
            if j == 4:
                cell.fill = FILL["assumption"]
            if j == 9:
                cell.fill = FILL["market data"]
    n = len(COUNTRIES) + 1
    notes = [
        "VAT standard: PwC Worldwide Tax Summaries (AA3, K1). Reduced: rates that may apply to cultural events or performances; dance parties are often at the standard rate: verify.",
        "Ticket platform: Eventbrite Professional package, per country (AA3, K2). goandance does not publish its commission.",
        "Music rights: Germany GEMA tariff, about 500 m2, 500-1,500 people; Austria AKM minimum fee for 401-500 people (AA3, K4). Empty = not found: the model uses the German value.",
        "Price level: Eurostat tec00120, 2024. Market median: only countries with 5+ events with a price (AA1, AB).",
    ]
    for k, t in enumerate(notes):
        c = wc.cell(row=n + 2 + k, column=1, value=t)
        c.alignment = Alignment(wrap_text=False)

    # Inputs -------------------------------------------------------------
    wi = wb.create_sheet("Inputs", 1)
    style_header(wi, 1, ["Input", "Unit", "Small", "Medium", "Large", "Basis", "Note"], [46, 14, 12, 12, 12, 13, 90])
    wi.cell(row=2, column=1, value="General inputs (one value for all scenarios: column Small)").font = BOLD
    G, S = {}, {}
    r = 3
    for key, label, unit, val, basis, note in GLOBAL:
        wi.cell(row=r, column=1, value=label)
        wi.cell(row=r, column=2, value=unit)
        c = wi.cell(row=r, column=3, value=val)
        c.fill = FILL[basis]
        if unit.startswith("%"):
            c.number_format = PCT
        elif unit.startswith("EUR"):
            c.number_format = EUR2 if isinstance(val, float) and val < 10 else EUR
        for col in (4, 5):
            wi.cell(row=r, column=col, value="")
        wi.cell(row=r, column=6, value=basis)
        wi.cell(row=r, column=7, value=note)
        for col in range(1, 8):
            wi.cell(row=r, column=col).border = THIN
        G[key] = f"Inputs!$C${r}"
        r += 1
    dv = DataValidation(type="list", formula1='"' + ",".join(c[0] for c in COUNTRIES) + '"', allow_blank=False)
    wi.add_data_validation(dv)
    dv.add(G["country"].split("!")[1].replace("$", ""))
    yn = DataValidation(type="list", formula1='"yes,no"', allow_blank=False)
    wi.add_data_validation(yn)
    for k in ("reduced_vat", "scale_prices", "floor_needed"):
        yn.add(G[k].split("!")[1].replace("$", ""))

    r += 1
    wi.cell(row=r, column=1, value="Scenario inputs").font = BOLD
    r += 1
    for key, label, unit, vals, basis, note in SCEN:
        wi.cell(row=r, column=1, value=label)
        wi.cell(row=r, column=2, value=unit)
        for j, v in enumerate(vals):
            c = wi.cell(row=r, column=3 + j, value=v)
            c.fill = FILL[basis]
            c.number_format = PCT if unit == "%" else (EUR if unit.startswith("EUR") else "#,##0.##")
        wi.cell(row=r, column=6, value=basis)
        wi.cell(row=r, column=7, value=note)
        for col in range(1, 8):
            wi.cell(row=r, column=col).border = THIN
        S[key] = r
        r += 1

    r += 1
    wi.cell(row=r, column=1, value="Taken from the sheet Country for the chosen country (formulas)").font = BOLD
    r += 1
    country = G["country"]
    derived = [
        ("vat", "VAT rate used", "%", f'IF({G["reduced_vat"]}="yes",INDEX(Country!$D$2:$D${n},MATCH({country},Country!$A$2:$A${n},0)),INDEX(Country!$C$2:$C${n},MATCH({country},Country!$A$2:$A${n},0)))'),
        ("plat_pct", "Ticket platform fee %", "%", f"INDEX(Country!$E$2:$E${n},MATCH({country},Country!$A$2:$A${n},0))"),
        ("plat_fixed", "Ticket platform fee, fixed per ticket", "EUR", f"INDEX(Country!$F$2:$F${n},MATCH({country},Country!$A$2:$A${n},0))"),
        ("music", "Music rights per night", "EUR", f'IF(INDEX(Country!$G$2:$G${n},MATCH({country},Country!$A$2:$A${n},0))="",Country!$G$3,INDEX(Country!$G$2:$G${n},MATCH({country},Country!$A$2:$A${n},0)))'),
        ("factor", "Local price factor (country / Germany)", "x", f'IF({G["scale_prices"]}="yes",INDEX(Country!$H$2:$H${n},MATCH({country},Country!$A$2:$A${n},0))/{BASE_PLI},1)'),
        ("market_pass", "Market median full pass in this country (reference)", "EUR", f'IFERROR(IF(INDEX(Country!$I$2:$I${n},MATCH({country},Country!$A$2:$A${n},0))="","fewer than 5 events",INDEX(Country!$I$2:$I${n},MATCH({country},Country!$A$2:$A${n},0))),"")'),
    ]
    for key, label, unit, formula in derived:
        wi.cell(row=r, column=1, value=label)
        wi.cell(row=r, column=2, value=unit)
        c = wi.cell(row=r, column=3, value="=" + formula)
        c.fill = FILL["formula"]
        c.number_format = PCT if unit == "%" else (EUR2 if unit == "EUR" else "0.000")
        wi.cell(row=r, column=6, value="formula")
        G[key] = f"Inputs!$C${r}"
        r += 1
    wi.freeze_panes = "C2"

    def s(key, col):
        return f"Inputs!${col}${S[key]}"

    # Model --------------------------------------------------------------
    wm = wb.create_sheet("Model", 2)
    style_header(wm, 1, ["Line", "How it is computed", "Small", "Medium", "Large"], [34, 70, 14, 14, 14])
    M = {}
    row = [2]

    def section(title):
        c = wm.cell(row=row[0], column=1, value=title)
        c.font = WHITE
        c.fill = FILL["band"]
        for col in range(2, 6):
            wm.cell(row=row[0], column=col).fill = FILL["band"]
        row[0] += 1

    def line(key, label, how, f, bold=False):
        wm.cell(row=row[0], column=1, value=label).font = Font(bold=bold)
        wm.cell(row=row[0], column=2, value=how).alignment = Alignment(wrap_text=False)
        for col in COLS:
            c = wm.cell(row=row[0], column=COLS.index(col) + 3, value="=" + f(col))
            c.number_format = EUR
            c.fill = FILL["formula"]
            if bold:
                c.font = BOLD
        M[key] = row[0]
        row[0] += 1

    def m(key, col):
        return f"{col}{M[key]}"

    g = G
    section("Revenue")
    line("gross", "Ticket sales, gross (incl. VAT)", "dancers x (full share x full pass + rest x party pass) x (1 - discount)",
         lambda c: f'{s("dancers",c)}*({s("full_share",c)}*{g["pass_price"]}+(1-{s("full_share",c)})*{g["party_price"]})*(1-{g["discount"]})')
    line("net", "Ticket sales, net of VAT", "gross / (1 + VAT)", lambda c: f'{m("gross",c)}/(1+{g["vat"]})')
    line("hotel", "Hotel commission", "dancers x hotel share x price per night x nights x commission",
         lambda c: f'{s("dancers",c)}*{g["hotel_share"]}*{g["hotel_pppn"]}*{g["nights"]}*{g["hotel_comm"]}')
    line("sponsor", "Sponsors and stands", "input", lambda c: s("sponsor", c))
    line("bar", "Share of bar sales", "input", lambda c: s("bar", c))
    line("rev", "Total revenue", "sum", lambda c: f'SUM({c}{M["net"]}:{c}{M["bar"]})', bold=True)

    section("Artists and programme")
    first_prod = row[0]
    line("c_fee", "Artist fees", "teaching artists x fee", lambda c: f'{s("units",c)}*{g["artist_fee"]}')
    line("c_fly", "Artist flights", "artists x people x share flying x fare x 2 x (1 + extras)",
         lambda c: f'{s("units",c)}*{g["people_per_unit"]}*{g["fly_share"]}*{g["flight"]}*2*(1+{g["flight_extra"]})')
    line("c_hot", "Artist hotel", "artists x room per night x days x price factor", lambda c: f'{s("units",c)}*{g["artist_room"]}*{g["days"]}*{g["factor"]}')
    line("c_meal", "Artist meals", "artists x people x meals per day x days x price factor",
         lambda c: f'{s("units",c)}*{g["people_per_unit"]}*{g["per_diem"]}*{g["days"]}*{g["factor"]}')
    line("c_dj", "DJs", "DJs x (fee x nights + travel)", lambda c: f'{s("djs",c)}*({g["dj_fee"]}*{g["nights"]}+{g["dj_travel"]})')

    section("Venue and production")
    line("c_hall", "Party hall", "rent per day x days x price factor", lambda c: f'{s("hall_day",c)}*{g["days"]}*{g["factor"]}')
    line("c_rooms", "Workshop rooms", "rooms x rent per day x days x price factor", lambda c: f'{s("rooms",c)}*{g["room_day"]}*{g["days"]}*{g["factor"]}')
    line("c_floor", "Dance floor", "if needed: m2 x price per m2 x price factor",
         lambda c: f'IF({g["floor_needed"]}="yes",{s("floor_area",c)}*{g["floor_m2"]}*{g["factor"]},0)')
    line("c_sound", "Sound and light", "(party hall + rooms x room package) x days x price factor",
         lambda c: f'({s("sound_main",c)}+{s("rooms",c)}*{g["sound_room"]})*{g["days"]}*{g["factor"]}')
    line("c_music", "Music rights", "fee per night x size factor x nights", lambda c: f'{g["music"]}*{s("music_factor",c)}*{g["nights"]}')
    line("c_deco", "Decoration and stage", "input x price factor", lambda c: f'{s("deco",c)}*{g["factor"]}')

    section("Staff and services")
    line("c_sec", "Security", "staff x hours x rate x nights x price factor",
         lambda c: f'{s("guards",c)}*{g["hours_night"]}*{g["security_rate"]}*{g["nights"]}*{g["factor"]}')
    line("c_aid", "First aid", "staff x hours x rate x nights x price factor",
         lambda c: f'{s("aid_staff",c)}*{g["hours_night"]}*{g["aid_rate"]}*{g["nights"]}*{g["factor"]}')
    line("c_photo", "Photo and video", "people x day rate x days x price factor", lambda c: f'{s("photographers",c)}*{g["photo_day"]}*{g["days"]}*{g["factor"]}')
    line("c_vol", "Volunteers' meals", "volunteers x meals x days x price factor", lambda c: f'{s("volunteers",c)}*{g["meals_vol"]}*{g["days"]}*{g["factor"]}')
    line("c_band", "Wristbands and passes", "dancers x price per piece", lambda c: f'{s("dancers",c)}*{g["wristband"]}')
    line("c_ins", "Insurance", "input", lambda c: s("insurance", c))
    line("c_admin", "Website, software, admin", "input", lambda c: s("admin", c))
    line("c_legal", "Accounting and legal", "input", lambda c: s("legal", c))
    last_prod = row[0] - 1

    section("Fees, marketing, contingency")
    line("c_cont", "Contingency", "contingency % x production costs above",
         lambda c: f'{g["contingency"]}*(SUM({c}{first_prod}:{c}{last_prod}))')
    line("c_plat", "Ticket platform fees", "gross sales x fee % + dancers x fixed fee", lambda c: f'{m("gross",c)}*{g["plat_pct"]}+{s("dancers",c)}*{g["plat_fixed"]}')
    line("c_pay", "Card payment fees", "gross sales x fee % + dancers x fixed fee", lambda c: f'{m("gross",c)}*{g["pay_pct"]}+{s("dancers",c)}*{g["pay_fixed"]}')
    line("c_mkt", "Marketing", "marketing % x total revenue", lambda c: f'{g["marketing_pct"]}*{m("rev",c)}')
    line("cost", "Total costs", "sum of all cost lines",
         lambda c: f'SUM({c}{first_prod}:{c}{last_prod})+{m("c_cont",c)}+{m("c_plat",c)}+{m("c_pay",c)}+{m("c_mkt",c)}', bold=True)

    section("Result")
    line("profit", "Result (revenue - costs)", "", lambda c: f'{m("rev",c)}-{m("cost",c)}', bold=True)
    wm.freeze_panes = "C2"

    # Results ------------------------------------------------------------
    wr = wb.create_sheet("Results", 1)
    style_header(wr, 1, ["Measure", "How it is computed", "Small", "Medium", "Large"], [40, 72, 14, 14, 14])
    R = {}
    rr = [2]

    def res(key, label, how, f, fmt=EUR, bold=False):
        wr.cell(row=rr[0], column=1, value=label).font = Font(bold=bold)
        wr.cell(row=rr[0], column=2, value=how)
        for col in COLS:
            c = wr.cell(row=rr[0], column=COLS.index(col) + 3, value="=" + f(col))
            c.number_format = fmt
            c.fill = FILL["formula"]
            if bold:
                c.font = BOLD
        R[key] = rr[0]
        rr[0] += 1

    res("d", "Paying dancers", "input", lambda c: s("dancers", c), "#,##0")
    res("rev", "Total revenue", "Model", lambda c: f"Model!{c}{M['rev']}")
    res("cost", "Total costs", "Model", lambda c: f"Model!{c}{M['cost']}")
    res("profit", "Result", "revenue - costs", lambda c: f"Model!{c}{M['profit']}", bold=True)
    res("margin", "Margin", "result / revenue", lambda c: f"IF(Model!{c}{M['rev']}=0,0,Model!{c}{M['profit']}/Model!{c}{M['rev']})", PCT)
    res("cpd", "Cost per paying dancer", "costs / dancers", lambda c: f"Model!{c}{M['cost']}/{s('dancers',c)}")
    res("rpd", "Net revenue per dancer (tickets and hotel)", "(net ticket sales + hotel commission) / dancers",
        lambda c: f"(Model!{c}{M['net']}+Model!{c}{M['hotel']})/{s('dancers',c)}")
    res("vpd", "Variable cost per dancer", "fees, wristband, marketing and its contingency per dancer",
        lambda c: (f"(Model!{c}{M['c_plat']}+Model!{c}{M['c_pay']})/{s('dancers',c)}+{g['wristband']}*(1+{g['contingency']})"
                   f"+{g['marketing_pct']}*(Model!{c}{M['net']}+Model!{c}{M['hotel']})/{s('dancers',c)}"))
    res("fixed", "Fixed costs of this set-up", "all other costs, net of sponsors and bar",
        lambda c: (f"Model!{c}{M['cost']}-{s('dancers',c)}*{c}{R['vpd']}"
                   f"-Model!{c}{M['sponsor']}-Model!{c}{M['bar']}"))
    res("be_d", "Break-even dancers", "fixed costs / (net revenue - variable cost) per dancer",
        lambda c: f'IF({c}{R["rpd"]}-{c}{R["vpd"]}<=0,"never",{c}{R["fixed"]}/({c}{R["rpd"]}-{c}{R["vpd"]}))', "#,##0", bold=True)
    res("slope", "Result change per EUR of full pass price", "full-pass dancers x (1 - discount) x (net share - fees)",
        lambda c: (f"{s('dancers',c)}*{s('full_share',c)}*(1-{g['discount']})*((1-{g['marketing_pct']})/(1+{g['vat']})"
                   f"-{g['plat_pct']}-{g['pay_pct']})"), EUR2)
    res("be_p", "Break-even full pass price (listed)", "price - result / change per EUR",
        lambda c: f"{g['pass_price']}-{c}{R['profit']}/{c}{R['slope']}", EUR, bold=True)
    rr[0] += 1
    wr.cell(row=rr[0], column=1, value="Read with care: orange inputs on Inputs are placeholders. Replace them with real offers before using the result.").font = Font(italic=True)
    rr[0] += 1
    wr.cell(row=rr[0], column=1, value="Biggest cost lines (medium scenario)").font = BOLD
    rr[0] += 1
    cost_keys = ["c_fee", "c_fly", "c_hot", "c_meal", "c_dj", "c_hall", "c_rooms", "c_floor", "c_sound", "c_music", "c_deco",
                 "c_sec", "c_aid", "c_photo", "c_vol", "c_band", "c_ins", "c_admin", "c_legal", "c_cont", "c_plat", "c_pay", "c_mkt"]
    for k in cost_keys:
        wr.cell(row=rr[0], column=1, value=f"=Model!A{M[k]}")
        c = wr.cell(row=rr[0], column=4, value=f"=Model!D{M[k]}")
        c.number_format = EUR
        c2 = wr.cell(row=rr[0], column=5, value=f"=IF(Model!D{M['cost']}=0,0,Model!D{M[k]}/Model!D{M['cost']})")
        c2.number_format = PCT
        rr[0] += 1
    wr.cell(row=rr[0], column=4, value="EUR").font = Font(italic=True)
    wr.cell(row=rr[0], column=5, value="share of costs").font = Font(italic=True)

    # Sensitivity --------------------------------------------------------
    wsn = wb.create_sheet("Sensitivity", 3)
    style_header(wsn, 1, ["Line or input (medium scenario)", "Value", "Result if 20% lower", "Result if 20% higher", "Swing"], [40, 14, 18, 18, 14])
    medium_lines = py_model()[1][3]
    label_of = dict(zip(cost_keys, medium_lines.keys()))
    order = sorted(cost_keys, key=lambda k: -medium_lines[label_of[k]])
    rs = 2
    # price row
    wsn.cell(row=rs, column=1, value="Full pass price")
    wsn.cell(row=rs, column=2, value=f"={g['pass_price']}").number_format = EUR
    wsn.cell(row=rs, column=3, value=f"=Results!D{R['profit']}-0.2*{g['pass_price']}*Results!D{R['slope']}").number_format = EUR
    wsn.cell(row=rs, column=4, value=f"=Results!D{R['profit']}+0.2*{g['pass_price']}*Results!D{R['slope']}").number_format = EUR
    wsn.cell(row=rs, column=5, value=f"=D{rs}-C{rs}").number_format = EUR
    rs += 1
    wsn.cell(row=rs, column=1, value="Paying dancers (same set-up)")
    wsn.cell(row=rs, column=2, value=f"={s('dancers','D')}").number_format = "#,##0"
    wsn.cell(row=rs, column=3, value=f"=Results!D{R['profit']}-0.2*{s('dancers','D')}*(Results!D{R['rpd']}-Results!D{R['vpd']})").number_format = EUR
    wsn.cell(row=rs, column=4, value=f"=Results!D{R['profit']}+0.2*{s('dancers','D')}*(Results!D{R['rpd']}-Results!D{R['vpd']})").number_format = EUR
    wsn.cell(row=rs, column=5, value=f"=D{rs}-C{rs}").number_format = EUR
    rs += 1
    for k in order:
        wsn.cell(row=rs, column=1, value=f"=Model!A{M[k]}")
        wsn.cell(row=rs, column=2, value=f"=Model!D{M[k]}").number_format = EUR
        wsn.cell(row=rs, column=3, value=f"=Results!D{R['profit']}+0.2*Model!D{M[k]}").number_format = EUR
        wsn.cell(row=rs, column=4, value=f"=Results!D{R['profit']}-0.2*Model!D{M[k]}").number_format = EUR
        wsn.cell(row=rs, column=5, value=f"=C{rs}-D{rs}").number_format = EUR
        rs += 1
    wsn.cell(row=rs + 1, column=1, value="Cost lines: 20% lower cost = higher result. Lines are ordered by size with the default inputs; the swing updates with your inputs.").font = Font(italic=True)

    # Market -------------------------------------------------------------
    wk = wb.create_sheet("Market")
    style_header(wk, 1, ["Group", "Events", "Median full pass EUR", "Note"], [26, 10, 20, 70])
    rows = [("All, random sample", len(FULL), statistics.median(FULL), "Latin congresses and festivals of the next 12 months drawn at random (seed 20261021)")]
    rows += [(k, len(v), statistics.median(v), "all events with a price") for k, v in sorted(BY_DAYS.items())]
    rows += [(f"Country {k}", len(v), statistics.median(v), "") for k, v in sorted(BY_COUNTRY.items(), key=lambda x: -len(x[1])) if len(v) >= 5]
    rows += [("Party-only pass", len(PARTY), statistics.median(PARTY), "events that publish one")]
    for i, (a, b, c_, d) in enumerate(rows, 2):
        wk.cell(row=i, column=1, value=a)
        wk.cell(row=i, column=2, value=b)
        cc = wk.cell(row=i, column=3, value=round(c_, 1))
        cc.number_format = EUR
        cc.fill = FILL["market data"]
        wk.cell(row=i, column=4, value=d)
    wk.cell(row=len(rows) + 3, column=1, value="Source: data/prices/latin_events_detail.csv (rounds AA1 and AB, October 2026); prices as listed on the access day, tier mostly unknown.")

    # Sources ------------------------------------------------------------
    wsrc = wb.create_sheet("Sources")
    style_header(wsrc, 1, ["Item", "Country", "Value EUR", "Low", "High", "Unit", "VAT incl.", "Year", "Source", "Link", "Notes"],
                 [22, 8, 11, 9, 9, 30, 9, 7, 28, 50, 70])
    i = 2
    for b in csv.DictReader(open(BENCH, encoding="utf-8")):
        if "Seattle" in b["notes"] or "Four Points" in b["source_name"]:
            continue
        vals = [b["item"], b["country_iso2"], b["value_eur"], b["value_low_eur"], b["value_high_eur"], b["unit"], b["vat_included"], b["year"], b["source_name"], b["source_url"], b["notes"]]
        for j, v in enumerate(vals, 1):
            if j in (3, 4, 5) and v:
                try:
                    v = float(v)
                except ValueError:
                    pass
            wsrc.cell(row=i, column=j, value=v)
        i += 1
    wsrc.cell(row=i + 1, column=1, value="K11 artist fees are private lesson rates, not congress fees: not used in the model. US hotel terms row removed.").font = Font(italic=True)

    wb.save(OUT)
    return M, R


def recalc_with_libreoffice():
    """Store computed values in the file (LibreOffice does not recalculate Excel files on load by default)."""
    import shutil
    import subprocess
    import tempfile
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice not found: open the file and press Ctrl+Shift+F9 to recalculate.")
        return
    tmp = tempfile.mkdtemp()
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, OUT],
                   capture_output=True, check=False)
    done = os.path.join(tmp, os.path.basename(OUT))
    if os.path.exists(done):
        shutil.copy(done, OUT)
        print("Recalculated with LibreOffice.")


def main():
    build()
    recalc_with_libreoffice()
    names = ["Small", "Medium", "Large"]
    for name, (rev, cost, profit, lines) in zip(names, py_model()):
        top = sorted(lines.items(), key=lambda x: -x[1])[:5]
        print(f"{name}: revenue {rev:,.0f}  costs {cost:,.0f}  result {profit:,.0f}  | top costs: "
              + ", ".join(f"{k} {v:,.0f}" for k, v in top))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
